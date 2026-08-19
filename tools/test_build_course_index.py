import tempfile
import unittest
from pathlib import Path

from pypdf import PdfWriter
from pptx import Presentation

from tools.build_course_index import build_index, classify_material, extract_pages, split_page
from tools.render_pdf_page import render_page


class CourseIndexTests(unittest.TestCase):
    def test_classifies_known_eam_material_types(self) -> None:
        examples = {
            Path("Lecture/L04 Foundations.pdf"): "lecture",
            Path("H05.pdf"): "exercise",
            Path("H05_Solution.pdf"): "solution",
            Path("MockExam.pdf"): "exam",
            Path("EAM_L1-L4_Bilingual_Revision_Notes.pdf"): "revision-note",
            Path("Exercises/debugging.md"): "exercise",
        }

        for path, expected in examples.items():
            with self.subTest(path=path):
                self.assertEqual(expected, classify_material(path))

    def test_classifies_real_world_multilingual_course_layouts(self) -> None:
        examples = {
            Path("Lectures/lec3_DiscreteDynamics.pdf"): "lecture",
            Path("Course Material/Lecture Slides/Lecture IV.pdf"): "lecture",
            Path("2026 LA chapter 10.pdf"): "lecture",
            Path("Tutorials/Exercise Sheet 4.pdf"): "exercise",
            Path("作业/Assignment 2.pdf"): "exercise",
            Path("Übung/Blatt 03.pdf"): "exercise",
            Path("Exercise Solutions/Problem Set 3 - Lösungen.pdf"): "solution",
            Path("Mock Exam/ExerciseExam.pdf"): "exam",
            Path("Signal_Processing_ExamNotes_CN.pdf"): "revision-note",
        }

        for path, expected in examples.items():
            with self.subTest(path=path):
                self.assertEqual(expected, classify_material(path))

    def test_chunk_overlap_preserves_boundary_context(self) -> None:
        text = "one two three four five six seven eight nine ten"

        chunks = split_page(text, chunk_words=6, overlap_words=2)

        self.assertEqual(["one two three four five six", "five six seven eight nine ten"], chunks)

    def test_indexes_markdown_and_text_materials(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory)
            (source / "Lecture").mkdir()
            (source / "Lecture" / "grounding.md").write_text(
                "Grounded tutoring keeps citations to evidence.", encoding="utf-8"
            )
            (source / "notes.txt").write_text("Page one\fPage two", encoding="utf-8")

            index = build_index(source, chunk_words=20, overlap_words=4, course_id="demo")

        self.assertEqual(2, index["statistics"]["documentCount"])
        self.assertEqual(3, index["statistics"]["pageCount"])
        self.assertEqual("lecture", index["documents"][0]["materialType"])

    def test_extract_pages_preserves_explicit_text_boundaries(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "notes.md"
            path.write_text("First page\fSecond page", encoding="utf-8")

            self.assertEqual(["First page", "Second page"], extract_pages(path))

    def test_extracts_powerpoint_slides_as_page_cited_units(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "lecture.pptx"
            presentation = Presentation()
            first = presentation.slides.add_slide(presentation.slide_layouts[1])
            first.shapes.title.text = "Grounded retrieval"
            first.placeholders[1].text = "Evidence keeps AI explanations auditable."
            second = presentation.slides.add_slide(presentation.slide_layouts[1])
            second.shapes.title.text = "Knowledge check"
            presentation.save(path)

            pages = extract_pages(path)

            self.assertEqual(2, len(pages))
            self.assertIn("Grounded retrieval", pages[0])
            self.assertIn("Knowledge check", pages[1])

    def test_skips_sidecars_and_unreadable_documents_without_rejecting_course(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory)
            (source / "Lecture").mkdir()
            (source / "Lecture" / "valid.md").write_text(
                "A valid lecture remains searchable.", encoding="utf-8"
            )
            (source / "Lecture" / "._sidecar.pdf").write_bytes(b"not a pdf")
            (source / "Lecture" / "broken.pdf").write_bytes(b"not a pdf either")

            index = build_index(source, chunk_words=20, overlap_words=4, course_id="robust")

        self.assertEqual(1, index["statistics"]["documentCount"])
        self.assertEqual(1, index["statistics"]["skippedDocumentCount"])
        self.assertEqual("Lecture/broken.pdf", index["warnings"][0]["relativePath"])

    def test_renders_pdf_page_as_png_preview(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_directory:
            source = Path(temporary_directory) / "lecture.pdf"
            output = Path(temporary_directory) / "preview.png"
            writer = PdfWriter()
            writer.add_blank_page(width=612, height=792)
            with source.open("wb") as stream:
                writer.write(stream)

            render_page(source, output, page_number=1, target_width=600)

            self.assertTrue(output.exists())
            self.assertEqual(b"\x89PNG\r\n\x1a\n", output.read_bytes()[:8])


if __name__ == "__main__":
    unittest.main()

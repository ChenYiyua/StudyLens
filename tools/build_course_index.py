"""Build a page-cited local search index from PDF, Markdown, and text materials.

The generated JSON deliberately stores relative filenames rather than the source
folder. This keeps the index portable and avoids leaking a user's local path.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

from pypdf import PdfReader


COURSE_ID = "tum-inhn0017-eam"
COURSE_NAME = "Enterprise Architecture Management and Reference Models"
WORD_PATTERN = re.compile(r"\S+")
WHITESPACE_PATTERN = re.compile(r"\s+")
SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".md", ".txt"}


@dataclass(frozen=True)
class Chunk:
    id: str
    document_id: str
    title: str
    relative_path: str
    material_type: str
    page: int
    text: str


def chunk_to_json(chunk: Chunk) -> dict[str, object]:
    return {
        "id": chunk.id,
        "documentId": chunk.document_id,
        "title": chunk.title,
        "relativePath": chunk.relative_path,
        "materialType": chunk.material_type,
        "page": chunk.page,
        "text": chunk.text,
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--source",
        required=True,
        type=Path,
        help="Folder containing PDF, Markdown, or text course materials",
    )
    parser.add_argument("--output", required=True, type=Path, help="JSON index to create")
    parser.add_argument("--course-id", default=COURSE_ID, help="Stable URL-safe course identifier")
    parser.add_argument("--course-name", default=COURSE_NAME, help="Human-readable course name")
    parser.add_argument("--chunk-words", type=int, default=220)
    parser.add_argument("--overlap-words", type=int, default=40)
    return parser.parse_args()


def classify_material(relative_path: Path) -> str:
    name = relative_path.stem.casefold()
    path = relative_path.as_posix().casefold()
    searchable_name = re.sub(r"[_\-.]+", " ", name)
    searchable_path = re.sub(r"[_\-.\\/]+", " ", path)
    compact_path = re.sub(r"[^\w]+", "", path)

    if "revisionnotes" in compact_path or "examnotes" in compact_path or any(
        keyword in searchable_path
        for keyword in ("revision note", "exam note", "cheat sheet", "cheatsheet", "复习", "总结")
    ):
        return "revision-note"
    if any(
        keyword in searchable_path
        for keyword in ("solution", "solutions", "answer key", "model answer", "loesung", "lösung", "答案", "解答")
    ):
        return "solution"
    if any(keyword in compact_path for keyword in ("mockexam", "modelexam", "finalexam", "shorttest")) or any(
        keyword in searchable_path
        for keyword in ("mock exam", "model exam", "final exam", "retake", "skillcheck", "skill check", "short test", "klausur", "prüfung", "pruefung", "考试", "试卷")
    ) or re.search(r"(?:^|\s)exam(?:\s|$)", searchable_path):
        return "exam"
    if any(
        keyword in searchable_path
        for keyword in ("exercise", "exercises", "assignment", "assignments", "problem set", "worksheet", "homework", "tutorial", "tutorials", "uebung", "übung", "ubung", "aufgabe", "tutorium", "作业", "习题", "练习")
    ):
        return "exercise"
    if re.search(r"(?:^|\s)(?:ex|hw|h)\s*\d+", searchable_name):
        return "exercise"
    if any(
        keyword in searchable_path
        for keyword in ("lecture", "lectures", "lecture slides", "chapter", "chapters", "vorlesung", "folien", "skript", "课件", "讲义")
    ):
        return "lecture"
    if re.search(r"(?:^|\s)(?:lec|lecture|l)\s*\d+", searchable_name):
        return "lecture"
    if re.match(r"^\d+[a-z]?\s+", searchable_name):
        return "lecture"
    if "case" in name or "archihotel" in name or "studyproject" in name or "c2f1g2" in path:
        return "case-study"
    return "supplement"


def stable_id(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()[:16]


def normalize_text(value: str) -> str:
    return WHITESPACE_PATTERN.sub(" ", value).strip()


def split_page(text: str, chunk_words: int, overlap_words: int) -> list[str]:
    words = WORD_PATTERN.findall(text)
    if not words:
        return []

    chunks: list[str] = []
    step = chunk_words - overlap_words
    for start in range(0, len(words), step):
        current = words[start : start + chunk_words]
        if not current:
            break
        chunks.append(" ".join(current))
        if start + chunk_words >= len(words):
            break
    return chunks


def extract_pages(material_path: Path) -> list[str]:
    """Return page-like text units while keeping PDF page numbers auditable."""
    if material_path.suffix.lower() == ".pdf":
        return [page.extract_text() or "" for page in PdfReader(str(material_path)).pages]
    if material_path.suffix.lower() == ".pptx":
        from pptx import Presentation

        presentation = Presentation(str(material_path))
        slides: list[str] = []
        for slide in presentation.slides:
            slide_text = [getattr(shape, "text", "") for shape in slide.shapes]
            slides.append("\n".join(text for text in slide_text if text))
        return slides

    # Form-feed is a portable explicit page boundary for plain-text fixtures.
    return material_path.read_text(encoding="utf-8").split("\f")


def build_index(
    source: Path,
    chunk_words: int,
    overlap_words: int,
    course_id: str = COURSE_ID,
    course_name: str = COURSE_NAME,
) -> dict[str, object]:
    if not source.is_dir():
        raise ValueError(f"Course folder does not exist: {source}")
    if overlap_words < 0 or chunk_words <= overlap_words:
        raise ValueError("chunk-words must be greater than overlap-words")

    material_paths = sorted(
        (
            path
            for path in source.rglob("*")
            if path.suffix.lower() in SUPPORTED_SUFFIXES
            and not path.name.startswith("._")
            and "__MACOSX" not in path.parts
            and not path.name.startswith("~$")
        ),
        key=lambda path: path.as_posix().lower(),
    )
    if not material_paths:
        raise ValueError(f"No supported course materials found under: {source}")

    chunks: list[Chunk] = []
    documents: list[dict[str, object]] = []
    total_pages = 0
    empty_pages = 0
    skipped_documents: list[dict[str, str]] = []

    for material_path in material_paths:
        relative_path = material_path.relative_to(source)
        portable_path = relative_path.as_posix()
        document_id = stable_id(portable_path.lower())
        material_type = classify_material(relative_path)
        try:
            pages = extract_pages(material_path)
        except Exception as exception:
            skipped_documents.append(
                {
                    "relativePath": portable_path,
                    "message": f"{type(exception).__name__}: {exception}",
                }
            )
            continue
        document_chunks = 0

        for page_number, raw_text in enumerate(pages, start=1):
            total_pages += 1
            text = normalize_text(raw_text)
            if not text:
                empty_pages += 1
                continue

            for chunk_number, chunk_text in enumerate(
                split_page(text, chunk_words, overlap_words), start=1
            ):
                chunks.append(
                    Chunk(
                        id=stable_id(f"{portable_path}:{page_number}:{chunk_number}"),
                        document_id=document_id,
                        title=material_path.stem,
                        relative_path=portable_path,
                        material_type=material_type,
                        page=page_number,
                        text=chunk_text,
                    )
                )
                document_chunks += 1

        documents.append(
            {
                "id": document_id,
                "title": material_path.stem,
                "relativePath": portable_path,
                "materialType": material_type,
                "pageCount": len(pages),
                "chunkCount": document_chunks,
            }
        )

    return {
        "schemaVersion": 1,
        "course": {"id": course_id, "name": course_name},
        "generatedAtUtc": datetime.now(timezone.utc).isoformat(),
        "statistics": {
            "documentCount": len(documents),
            "pageCount": total_pages,
            "emptyPageCount": empty_pages,
            "chunkCount": len(chunks),
            "skippedDocumentCount": len(skipped_documents),
        },
        "documents": documents,
        "chunks": [chunk_to_json(chunk) for chunk in chunks],
        "warnings": skipped_documents,
    }


def main() -> None:
    args = parse_args()
    index = build_index(
        args.source.resolve(),
        args.chunk_words,
        args.overlap_words,
        args.course_id,
        args.course_name,
    )
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(index, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )

    stats = index["statistics"]
    print(
        "Indexed "
        f"{stats['documentCount']} documents, {stats['pageCount']} pages and "
        f"{stats['chunkCount']} searchable chunks."
    )
    print(f"Skipped {stats['emptyPageCount']} pages without extractable text.")
    if stats["skippedDocumentCount"]:
        print(f"Skipped {stats['skippedDocumentCount']} unreadable documents.")
    print(f"Output: {args.output.resolve()}")


if __name__ == "__main__":
    main()
